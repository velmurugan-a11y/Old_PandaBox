"""Build docs/PandaBox_Command_Test_Results_V2901.xlsx from the tracker and a tracker_suite.py run.

    python tools/make_results_xlsx.py <tracker.xlsx> <final_run.json> <out.xlsx>
"""
import json
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

tracker_path, run_path, out_path = sys.argv[1:4]
HAPPY = sys.argv[4] if len(sys.argv) > 4 else ""
PRODUCT = sys.argv[5] if len(sys.argv) > 5 else ""
run = json.load(open(run_path))

# cases the bench cannot exercise, or where this firmware deliberately follows V2.89 over the V3.01 tracker
NOT_RUN = {
    "PB-001": "Visual: LD8 boot LED / power switch (not observable from the bench PC)",
    "PB-055": "Needs the ticket printer on the meter",
    "PB-088": "1-hour stress not run; covered in part by 40 Start/Stop cycles + 6 full suite rounds",
    "PB-092": "BLE range walk test (physical)",
    "PB-100": "Visual: BLE LED blink during data traffic",
    "PB-102": "Phone lock / app background (needs the FleetPanda app)",
    "PB-103": "Phone lock / app background (needs the FleetPanda app)",
    "PB-104": "App-side 'Invalid Data' check above 300 gal/min (needs the app)",
    "PB-105": "Negative pulse rotation (simulator pulser cannot run backwards)",
    "PB-107": "Power removed mid-delivery (BoxReset covered in PB-041)",
    "PB-095": "Double meter needs a second meter on Port 1 (J1); the bench has one USB-RS232 adapter, on J2. "
              "Both products were validated one at a time on J2 (tools/product_suite.py)",
    "PB-096": "Double meter wired crosswise needs two adapters (wrong-node handling checked in product_suite.py)",
    "PB-126": "Needs two LCR.iQ meters (bench: LCR-II node 1 + LCR.iQ node 2)",
    "PB-127": "Net gallons need temperature compensation (simulator reports Net #18 = 0)",
    "PB-128": "Auto-stop at a NET preset is not modelled by the simulator (PresetNet accepted, PB-068)",
}
for n in list(range(89, 92)) + [93, 94] + list(range(109, 126)) + [129]:
    NOT_RUN.setdefault(f"PB-{n:03d}", "FleetPanda phone-app flow (needs the app on a phone)")

EXTRA = {
    "PB-002": ("PASS", "Scan finds 'PandaBrainBLE' (name ends with BLE) after every boot; LED not observed"),
    "PB-003": ("PASS", "Tester connects and runs every command; LED not observed"),
    "PB-008": ("DIFFERS", "Replies 'LxBoxStorage 0,0,25600,' (empty storage for an unknown meter) rather than an error"),
    "PB-011": ("DIFFERS", "GetData n,1 returns the oldest stored record (or an empty LxGetDataTs n,) at any time; "
                          "V3.01 answered 'LxGetData 1' before Start"),
    "PB-014": ("DIFFERS", "GetData n,0 always returns the live meter values (storage is not involved)"),
    "PB-072": ("DIFFERS", "V2.89 format 'LxBoxTime <unix>' (golden capture); no ',0,0' GPS suffix as in V3.01"),
    "PB-052": ("DIFFERS", "LxPause 0: V2.89 forwards Pause and a real LCR accepts it with no delivery open "
                          "(rc 0, devSt 0x21, golden capture 19:12:19); V3.01 blocks it in firmware (LxPause 1)"),
    "PB-097": ("PARTIAL", "Only J2 is wired; a J1 with no meter is reported offline while J2 keeps working "
                          "(PORT1 section of product_suite.py); unplug/replug verified in PB-106"),
    "PB-098": ("PARTIAL", "See PB-097 / PB-106"),
}

wb_in = openpyxl.load_workbook(tracker_path, data_only=True)
ws_in = wb_in["LCR_BLE-Test_V3"]
cases = OrderedDict()
for r in ws_in.iter_rows(min_row=12, values_only=True):
    cid = r[3] if len(r) > 3 else None
    if not cid or not str(cid).startswith("PB-"):
        continue
    steps = str(r[6] or "")
    m = re.findall(r"send:?\s*\n?\s*([A-Za-z][^\n]*)", steps)
    cases[cid] = {"feature": str(r[1] or "").replace("\n", " ").strip(),
                  "cmd": "; ".join(m), "expect": re.sub(r"\s+", " ", str(r[8] or ""))[:400]}

by_id = OrderedDict()
for x in run:
    by_id.setdefault(x["id"], []).append(x)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Results"
hdr = ["Test case", "Feature", "Tracker command", "Tracker expected (V3.01)", "V2.901 result",
       "Observed on the bench (BLE -> box -> RS232 -> LCR simulator)", "Notes"]
ws.append(hdr)
fills = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "DIFFERS": "FFEB9C", "PARTIAL": "FFEB9C", "NOT RUN": "EDEDED"}
counts = {}
for cid, c in cases.items():
    checks = by_id.get(cid, [])
    if checks:
        status = "PASS" if all(k["ok"] for k in checks) else "FAIL"
        observed = "\n".join(f"{k['command']} -> {' | '.join(l for l in k['response'] if l)}" for k in checks)
        notes = "; ".join(sorted({k["note"] for k in checks if k.get("note")}))
    elif cid in EXTRA:
        status, notes = EXTRA[cid]
        observed = ""
    else:
        status, observed, notes = "NOT RUN", "", NOT_RUN.get(cid, "")
    if cid in EXTRA and checks:
        notes = EXTRA[cid][1]
        if status == "PASS" and EXTRA[cid][0] == "DIFFERS":
            status = "DIFFERS"
    counts[status] = counts.get(status, 0) + 1
    ws.append([cid, c["feature"], c["cmd"], c["expect"], status, observed, notes])
    ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor=fills[status])

widths = [10, 20, 30, 60, 12, 80, 60]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=1):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="305496")
ws.freeze_panes = "B2"

log = wb.create_sheet("Check log")
log.append(["#", "Case", "Command", "Reply", "OK", "Why", "Note", "ms"])
for i, x in enumerate(run, 1):
    log.append([i, x["id"], x["command"], " | ".join(l for l in x["response"] if l),
                "PASS" if x["ok"] else "FAIL", x.get("why", ""), x.get("note", ""), x.get("ms")])
for i, w in enumerate([5, 9, 36, 90, 7, 30, 60, 8], 1):
    log.column_dimensions[get_column_letter(i)].width = w
for cell in log[1]:
    cell.font = Font(bold=True)
log.freeze_panes = "A2"

summ = wb.create_sheet("Summary", 0)
summ.append(["PandaBox V2.901 (Project 05, V2.89-compatible) - tracker results on the bench"])
summ["A1"].font = Font(bold=True, size=13)
summ.append([])
summ.append(["Setup", "PC pandabox-tester (BLE) -> PandaBox (GD32F305, bench build) -> RS232 Port 2 (J2) -> USB-RS232 COM7 -> "
                      "LCR simulator (LCR-II node 1, LCR.iQ node 2). Port 1 (J1): nothing connected."])
summ.append(["Automated checks", f"{sum(x['ok'] for x in run)}/{len(run)} passed (tools/tracker_suite.py)"])
summ.append(["Tester happy_flow.csv", HAPPY])
summ.append(["Product-wise suite", PRODUCT])
summ.append([])
summ.append(["Tracker cases by result", ""])
for k in ("PASS", "DIFFERS", "PARTIAL", "FAIL", "NOT RUN"):
    summ.append([k, counts.get(k, 0)])
summ.column_dimensions["A"].width = 26
summ.column_dimensions["B"].width = 110
wb.save(out_path)
print(out_path, counts)
